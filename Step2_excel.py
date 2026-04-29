import openpyxl
import os
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

output_folder = "output"
excel_file = os.path.join(output_folder, "exchange.xlsx")

header = ["ID", "BASE", "TARGET", "RATE", "CREATED_AT"]

def save_to_excel(data: dict) -> str:
    os.makedirs(output_folder, exist_ok=True)

    if os.path.exists(excel_file):
        workbook = openpyxl.load_workbook(excel_file)
        sheet = workbook.active
    else:
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.append(header)

    # ---- Header Style ----
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(fill_type="solid", start_color="4F81BD")
    header_alignment = Alignment(horizontal="center", vertical="center")
    header_border = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin")
    )

    for cell in sheet[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment
        cell.border = header_border

    # ---- Column width ----
    col_widths = [6, 10, 10, 12, 22]
    for i, width in enumerate(col_widths, start=1):
        sheet.column_dimensions[get_column_letter(i)].width = width

    next_id = sheet.max_row

    row = [
        next_id,
        data["base"],
        data["target"],
        data["rate"],
        data["created_at"]
    ]

    sheet.append(row)

    # Center data
    align = Alignment(horizontal="center", vertical="center")
    for cell in sheet[sheet.max_row]:
        cell.alignment = align

    workbook.save(excel_file)
    return excel_file


if __name__ == "__main__":
    from Step3_db import save_db
    from Step1_scrap import scrap_data

    data = scrap_data()
    record = save_db(data)

    file_path = save_to_excel(record)

    print(f"Saved to Excel: {file_path}")